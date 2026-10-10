import io
import json
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth.models import Group, Permission, User
from django.core import serializers
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TransactionTestCase, override_settings
from rest_framework.test import APIClient

from .data_backup import SCHEMA, apply_import, snapshot, validate_records
from .models import Animal, Story


class DataBackupTests(TransactionTestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.settings_override = override_settings(MEDIA_ROOT=self.media.name)
        self.settings_override.enable()
        self.addCleanup(self.media.cleanup)
        self.addCleanup(self.settings_override.disable)
        self.admin = User.objects.create_user(username="backup-admin", password="secret", is_staff=True)
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        self.animal = Animal.objects.create(id="A001", name="Bantay", image="animals/dog.jpg")
        target = Path(self.media.name) / "animals/dog.jpg"
        target.parent.mkdir()
        target.write_bytes(b"original-image")

    def document(self, **changes):
        record = json.loads(serializers.serialize("json", [self.animal]))[0]
        record["fields"].update(changes)
        return {"schema": SCHEMA, "records": [record]}

    def upload(self, document=None, *, raw=None, name="backup.json"):
        return SimpleUploadedFile(name, raw if raw is not None else json.dumps(document or self.document()).encode())

    def preview(self, document=None, **kwargs):
        return self.client.post("/admin_api/settings/import/preview/", {"file": self.upload(document, **kwargs)}, format="multipart")

    def confirm(self, preview, document=None, **kwargs):
        return self.client.post("/admin_api/settings/import/confirm/", {
            "file": self.upload(document, **kwargs), "preview_token": preview.data["preview_token"], "confirm": "true",
        }, format="multipart")

    def export(self, format_name):
        response = self.client.get("/admin_api/settings/export/", {"file_format": format_name})
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        content = b"".join(response.streaming_content)
        response.close()
        return content

    def test_preview_is_read_only_and_confirmation_updates_record(self):
        document = self.document(name="Updated Bantay")
        preview = self.preview(document)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["update"], 1)
        self.animal.refresh_from_db()
        self.assertEqual(self.animal.name, "Bantay")
        result = self.confirm(preview, document)
        self.assertEqual(result.status_code, 200, result.data)
        self.animal.refresh_from_db()
        self.assertEqual(self.animal.name, "Updated Bantay")

    def test_all_formats_round_trip_and_preserve_relationships(self):
        group = Group.objects.create(name="Shelter staff")
        permission = Permission.objects.first()
        group.permissions.add(permission)
        self.admin.groups.add(group)
        self.admin.user_permissions.add(permission)
        Story.objects.create(animal=self.animal, title="Rescued", date="2026-01-01", content="Safe now")
        for format_name in ("json", "xlsx", "zip"):
            with self.subTest(format_name=format_name):
                raw = self.export(format_name)
                preview = self.preview(raw=raw, name=f"backup.{format_name}")
                self.assertEqual(preview.status_code, 200, preview.data)
                result = self.confirm(preview, raw=raw, name=f"backup.{format_name}")
                self.assertEqual(result.status_code, 200, result.data)
                self.assertEqual(list(self.admin.groups.values_list("pk", flat=True)), [group.pk])
                self.assertEqual(list(group.permissions.values_list("pk", flat=True)), [permission.pk])
                self.assertEqual(Story.objects.get().animal_id, self.animal.pk)

    def test_zip_restores_uploads_and_keeps_unrelated_files_and_records(self):
        raw = self.export("zip")
        target = Path(self.media.name) / "animals/dog.jpg"
        target.write_bytes(b"changed")
        unrelated = Path(self.media.name) / "keep.txt"
        unrelated.write_text("keep")
        Animal.objects.create(id="A002", name="Keep me")
        preview = self.preview(raw=raw, name="backup.zip")
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["replace_files"], 1)
        self.assertEqual(target.read_bytes(), b"changed")
        result = self.confirm(preview, raw=raw, name="backup.zip")
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(target.read_bytes(), b"original-image")
        self.assertTrue(unrelated.exists())
        self.assertTrue(Animal.objects.filter(pk="A002").exists())

    def test_confirmation_requires_preview_exact_file_user_and_consent(self):
        document = self.document(name="Update")
        preview = self.preview(document)
        response = self.client.post("/admin_api/settings/import/confirm/", {
            "file": self.upload(document), "confirm": "true",
        }, format="multipart")
        self.assertEqual(response.status_code, 400)
        response = self.client.post("/admin_api/settings/import/confirm/", {
            "file": self.upload(document), "preview_token": preview.data["preview_token"],
        }, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.confirm(preview, self.document(name="Different")).status_code, 400)
        other = User.objects.create_user(username="other-admin", is_staff=True)
        self.client.force_authenticate(other)
        self.assertEqual(self.confirm(preview, document).status_code, 400)
        self.animal.refresh_from_db()
        self.assertEqual(self.animal.name, "Bantay")

    def test_expired_and_stale_previews_rejected(self):
        document = self.document()
        preview = self.preview(document)
        with patch("django.core.signing.time.time", return_value=9999999999):
            self.assertEqual(self.confirm(preview, document).status_code, 400)
        self.animal.name = "Changed meanwhile"
        self.animal.save()
        self.assertEqual(self.confirm(preview, document).status_code, 409)

    def test_non_admins_cannot_export_preview_or_import(self):
        regular = User.objects.create_user(username="regular")
        for user in (None, regular):
            self.client.force_authenticate(user)
            self.assertIn(self.client.get("/admin_api/settings/export/").status_code, (401, 403))
            self.assertIn(self.preview().status_code, (401, 403))
            response = self.client.post("/admin_api/settings/import/confirm/", {}, format="multipart")
            self.assertIn(response.status_code, (401, 403))

    def test_rejects_invalid_fields_unknown_models_duplicates_and_broken_relations(self):
        invalids = [self.document(age="invalid"), self.document(name="x" * 101),
                    self.document(age=1.5), self.document(age=True), self.document(name={"bad": "value"})]
        missing = self.document()
        del missing["records"][0]["fields"]["name"]
        invalids.append(missing)
        duplicate = self.document()
        duplicate["records"] *= 2
        invalids.append(duplicate)
        unknown = self.document()
        unknown["records"][0]["model"] = "sessions.session"
        invalids.append(unknown)
        story = Story.objects.create(animal=self.animal, title="Story", date="2026-01-01", content="Content")
        story.refresh_from_db()
        record = json.loads(serializers.serialize("json", [story]))[0]
        record["fields"]["animal"] = "MISSING"
        invalids.append({"schema": SCHEMA, "records": [record]})
        for document in invalids:
            with self.subTest(document=document):
                response = self.preview(document)
                self.assertEqual(response.status_code, 400, response.data)
                self.assertTrue(response.data["errors"])
        self.animal.refresh_from_db()
        self.assertEqual(self.animal.name, "Bantay")

    def test_unique_conflict_rolls_back_entire_import(self):
        group = Group.objects.create(name="Existing")
        document = self.document(name="Must roll back")
        document["records"].append({"model": "auth.group", "pk": group.pk + 1, "fields": {"name": "Existing", "permissions": []}})
        response = self.preview(document)
        self.assertEqual(response.status_code, 400, response.data)
        self.animal.refresh_from_db()
        self.assertEqual(self.animal.name, "Bantay")

    def test_missing_files_warn_and_unsafe_paths_fail(self):
        preview = self.preview(self.document(image="animals/missing.jpg"))
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertTrue(preview.data["warnings"])
        for path in ("../outside.jpg", "C:/outside.jpg", "animals/CON.txt", "animals\\outside.jpg"):
            self.assertEqual(self.preview(self.document(image=path)).status_code, 400)

    def test_profile_placeholders_do_not_hide_missing_uploaded_files(self):
        from accounts.models import UserProfile as AccountProfile
        from user_api.models import UserProfile

        for model in (AccountProfile, UserProfile):
            with self.subTest(model=model._meta.label_lower):
                profile = model.objects.create(user=self.admin)
                records = json.loads(serializers.serialize("json", [profile]))
                document = {"schema": SCHEMA, "records": records}
                response = self.preview(document)
                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual(response.data["warnings"], [])

                records[0]["fields"]["avatar"] = "avatars/missing.jpg"
                records[0]["fields"]["cover_photo"] = "covers/missing.jpg"
                response = self.preview(document)
                self.assertEqual(response.status_code, 200, response.data)
                self.assertIn("2 referenced file(s)", response.data["warnings"][0])
                self.assertIn("avatars/missing.jpg", response.data["warnings"][0])
                self.assertIn("covers/missing.jpg", response.data["warnings"][0])

        response = self.preview(self.document(image="default.jpg"))
        self.assertIn("default.jpg", response.data["warnings"][0])

    def test_rejects_malformed_and_unsafe_archives(self):
        self.assertEqual(self.preview(raw=b"not-json").status_code, 400)
        self.assertEqual(self.preview(raw=b"not-excel", name="backup.xlsx").status_code, 400)
        for path in ("../escape.txt", "media/../../escape.txt", "unexpected.txt"):
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, "w") as archive:
                archive.writestr("records.json", json.dumps(self.document()))
                archive.writestr(path, b"unsafe")
            self.assertEqual(self.preview(raw=stream.getvalue(), name="backup.zip").status_code, 400)

    def test_file_failure_rolls_back_records_and_previous_uploads(self):
        document = self.document(name="Must roll back")
        objects = validate_records(document["records"], self.admin)
        files = {"animals/dog.jpg": b"replace", "animals/new.jpg": b"new"}
        import os
        real_replace = os.replace
        calls = []

        def fail_second(source, target):
            calls.append(target)
            if len(calls) == 2:
                raise OSError("Disk full")
            real_replace(source, target)

        with patch("admin_api.data_backup.os.replace", side_effect=fail_second):
            with self.assertRaises(OSError):
                apply_import(objects, files)
        self.animal.refresh_from_db()
        self.assertEqual(self.animal.name, "Bantay")
        self.assertEqual((Path(self.media.name) / "animals/dog.jpg").read_bytes(), b"original-image")
        self.assertFalse((Path(self.media.name) / "animals/new.jpg").exists())

    def test_account_lockout_is_rejected(self):
        document = snapshot()
        user = next(r for r in document["records"] if r["model"] == "auth.user")
        user["fields"]["is_staff"] = False
        self.assertEqual(self.preview(document).status_code, 400)

    def test_new_records_and_forward_references_restore(self):
        story = Story.objects.create(animal=self.animal, title="New story", date="2026-01-01", content="Safe")
        story.refresh_from_db()
        story_record = json.loads(serializers.serialize("json", [story]))[0]
        story_record["pk"] = 700
        story_record["fields"]["animal"] = "NEW"
        animal_record = self.document()["records"][0]
        animal_record["pk"] = "NEW"
        document = {"schema": SCHEMA, "records": [story_record, animal_record]}
        preview = self.preview(document)
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(preview.data["create"], 2)
        self.assertFalse(Animal.objects.filter(pk="NEW").exists())
        result = self.confirm(preview, document)
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(Story.objects.get(pk=700).animal_id, "NEW")
        next_story = Story.objects.create(animal=self.animal, title="Next", date="2026-01-01", content="Next")
        self.assertGreater(next_story.pk, 700)

    def test_excel_rejects_formulas_and_bad_schema(self):
        from openpyxl import load_workbook
        raw = self.export("xlsx")
        for invalid in ("formula", "schema"):
            workbook = load_workbook(io.BytesIO(raw))
            sheet_name = next(row[0] for row in workbook["Backup"].iter_rows(min_row=3, values_only=True) if row[1] == "admin_api.animal")
            sheet = workbook[sheet_name]
            if invalid == "formula":
                sheet["B2"] = '=HYPERLINK("https://example.com")'
            else:
                sheet["B1"] = "unknown_column"
            stream = io.BytesIO()
            workbook.save(stream)
            workbook.close()
            response = self.preview(raw=stream.getvalue(), name="backup.xlsx")
            self.assertEqual(response.status_code, 400, response.data)

    def test_zip_rejects_case_collisions_file_folder_conflicts_and_size_limit(self):
        for names in (("media/FILE.jpg", "media/file.jpg"), ("media/folder", "media/folder/file.jpg")):
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, "w") as archive:
                archive.writestr("records.json", json.dumps(self.document()))
                for name in names:
                    archive.writestr(name, b"test")
            self.assertEqual(self.preview(raw=stream.getvalue(), name="backup.zip").status_code, 400)
        raw = self.export("zip")
        with patch("admin_api.data_backup.MAX_EXPANDED", 1):
            self.assertEqual(self.preview(raw=raw, name="backup.zip").status_code, 400)
