"""Versioned, admin-only data interchange. Never includes sessions or API tokens."""
import hashlib
import io
import json
import os
import re
import shutil
import stat
import tempfile
import zipfile
from pathlib import Path

from django.apps import apps
from django.conf import settings
from django.core import serializers
from django.core.exceptions import ValidationError
from django.core.management.color import no_style
from django.db import IntegrityError, connection, models, transaction
from django.utils import timezone

SCHEMA = "aspin-backup-v1"
MAX_UPLOAD = 100 * 1024 * 1024
MAX_EXPANDED = 512 * 1024 * 1024
MAX_RECORD_BYTES = 50 * 1024 * 1024
MAX_RECORDS = 100000
MAX_FILES = 10000
APP_LABELS = {"accounts", "admin_api", "user_api", "chatbot"}
AUTH_MODELS = {"auth.user", "auth.group", "auth.permission", "contenttypes.contenttype"}


class BackupError(Exception):
    def __init__(self, errors):
        self.errors = errors if isinstance(errors, list) else [str(errors)]
        super().__init__(self.errors[0])


def backup_models():
    return {m._meta.label_lower: m for m in apps.get_models()
            if (m._meta.app_label in APP_LABELS or m._meta.label_lower in AUTH_MODELS)
            and m._meta.label_lower not in {
                'admin_api.assessmentemail', 'accounts.registrationverification', 'accounts.authemaillimit',
            }}


def model_fields(model):
    return [f for f in model._meta.local_fields + model._meta.local_many_to_many
            if f.serialize and not f.primary_key]


def check_value_type(field, value):
    if value is None:
        return
    if isinstance(field, models.BooleanField) and type(value) is not bool:
        raise BackupError(f"{field.name}: expected true or false.")
    if isinstance(field, models.IntegerField) and type(value) is not int:
        raise BackupError(f"{field.name}: expected a whole number.")
    if isinstance(field, (models.CharField, models.TextField, models.FileField)) and not isinstance(value, str):
        raise BackupError(f"{field.name}: expected text.")


def snapshot():
    records = []
    for model in backup_models().values():
        records.extend(json.loads(serializers.serialize("json", model.objects.order_by("pk"))))
    return {"schema": SCHEMA, "created_at": timezone.now().isoformat(), "records": records}


def json_bytes(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False).encode("utf-8")


def safe_path(name):
    # Also reject Windows aliases/alternate streams even on a Linux server.
    if not isinstance(name, str) or not name or "\\" in name:
        raise BackupError("Invalid file path in backup.")
    parts = name.split("/")
    for part in parts:
        if (part in {"", ".", ".."} or part.endswith((" ", "."))
                or re.search(r'[\x00-\x1f<>:"|?*]', part)
                or re.match(r"^(CON|PRN|AUX|NUL|COM[0-9]|LPT[0-9])(?:\.|$)", part, re.I)):
            raise BackupError(f"Unsafe file path: {name}")
    return name


def media_path(name):
    safe_path(name)
    root = Path(settings.MEDIA_ROOT).resolve()
    target = root.joinpath(*name.split("/"))
    if not target.resolve().is_relative_to(root):
        raise BackupError(f"File path is outside uploaded files: {name}")
    if any(p.is_symlink() for p in [target, *target.parents] if p != root and root in p.parents):
        raise BackupError(f"Symbolic links are not supported: {name}")
    if target.exists() and not target.is_file():
        raise BackupError(f"An uploaded file path is occupied by a directory: {name}")
    return target


def checked_archive(stream):
    archive = zipfile.ZipFile(stream)
    entries = archive.infolist()
    if len(entries) > MAX_FILES or sum(i.file_size for i in entries) > MAX_EXPANDED:
        archive.close()
        raise BackupError("Archive exceeds 10,000 entries or 512 MB uncompressed.")
    seen = set()
    for item in entries:
        name = safe_path(item.filename.rstrip("/") if item.is_dir() else item.filename)
        if name.casefold() in seen or stat.S_ISLNK(item.external_attr >> 16) or item.flag_bits & 1:
            archive.close()
            raise BackupError("Archive contains duplicate paths, symbolic links, or encrypted files.")
        seen.add(name.casefold())
    return archive


def excel_bytes(document):
    from openpyxl import Workbook
    workbook = Workbook()
    manifest = workbook.active
    manifest.title = "Backup"
    manifest.append([SCHEMA, document["created_at"]])
    manifest.append(["Sheet", "Model"])
    grouped = {label: [] for label in backup_models()}
    for record in document["records"]:
        grouped[record["model"]].append(record)
    for index, (label, rows) in enumerate(grouped.items(), 1):
        model = backup_models()[label]
        sheet = workbook.create_sheet(f"{index:02d}_{model._meta.model_name}"[:31])
        manifest.append([sheet.title, label])
        fields = [f.name for f in model_fields(model)]
        sheet.append(["pk", *fields])
        sheet.freeze_panes = "A2"
        for record in rows:
            values = [record["pk"], *[record["fields"][key] for key in fields]]
            encoded = [json.dumps(v, ensure_ascii=False, allow_nan=False) for v in values]
            if any(len(value) > 32767 for value in encoded):
                raise BackupError("A field exceeds Excel's cell limit. Export JSON or ZIP instead.")
            sheet.append(encoded)
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_excel(stream):
    from openpyxl import load_workbook
    with checked_archive(stream):
        pass
    stream.seek(0)
    workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
    try:
        if "Backup" not in workbook or workbook["Backup"].cell(1, 1).value != SCHEMA:
            raise BackupError("Use an Excel workbook exported from Settings (missing backup version).")
        records, seen = [], set()
        labels = backup_models()
        for sheet_name, label in workbook["Backup"].iter_rows(min_row=3, max_col=2, values_only=True):
            if label not in labels or label in seen or sheet_name not in workbook:
                raise BackupError("Invalid or duplicate model in Excel's Backup sheet.")
            seen.add(label)
            sheet = workbook[sheet_name]
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            headers = [cell.value for cell in next(rows)]
            expected = ["pk", *[f.name for f in model_fields(labels[label])]]
            if headers != expected:
                raise BackupError(f"{sheet_name}: columns do not match this application's schema.")
            for row_number, row in enumerate(rows, 2):
                if all(cell.value is None for cell in row):
                    continue
                if len(row) != len(headers):
                    raise BackupError(f"{sheet_name}, row {row_number}: incorrect column count.")
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        raise BackupError(f"{sheet_name}, {cell.coordinate}: formulas are not allowed.")
                    try:
                        # Exported cells use JSON literals to preserve types and long IDs.
                        values.append(json.loads(cell.value) if isinstance(cell.value, str) else cell.value)
                    except (ValueError, TypeError):
                        raise BackupError(f"{sheet_name}, {cell.coordinate}: use a JSON value, such as a quoted string.")
                records.append({"model": label, "pk": values[0], "fields": dict(zip(headers[1:], values[1:]))})
                if len(records) > MAX_RECORDS:
                    raise BackupError("Import exceeds 100,000 records.")
        if set(workbook.sheetnames) != {"Backup", *[r[0] for r in workbook["Backup"].iter_rows(min_row=3, max_col=2, values_only=True)]}:
            raise BackupError("Workbook contains sheets not listed in the Backup sheet.")
        return {"schema": SCHEMA, "records": records}
    finally:
        workbook.close()


def parse_upload(upload):
    if not upload or upload.size > MAX_UPLOAD:
        raise BackupError("Select a backup file up to 100 MB.")
    extension = Path(upload.name).suffix.lower()
    if extension not in {".json", ".xlsx", ".zip"}:
        raise BackupError("Supported formats are JSON, Excel (.xlsx), and ZIP.")
    upload.seek(0)
    raw = upload.read(MAX_UPLOAD + 1)
    files = {}
    try:
        if extension == ".xlsx":
            document = read_excel(io.BytesIO(raw))
        elif extension == ".zip":
            with checked_archive(io.BytesIO(raw)) as archive:
                if "records.json" not in archive.namelist():
                    raise BackupError("ZIP must contain records.json and optional media/ files.")
                if archive.getinfo("records.json").file_size > MAX_RECORD_BYTES:
                    raise BackupError("records.json exceeds 50 MB.")
                document = json.loads(archive.read("records.json"))
                for item in archive.infolist():
                    if item.is_dir():
                        continue
                    if item.filename == "records.json":
                        continue
                    if not item.filename.startswith("media/"):
                        raise BackupError(f"Unexpected ZIP entry: {item.filename}")
                    name = safe_path(item.filename[6:])
                    media_path(name)
                    files[name] = archive.read(item)
                paths = {name.casefold() for name in files}
                for name in paths:
                    if any(str(parent).replace("\\", "/") in paths for parent in Path(name).parents):
                        raise BackupError(f"A ZIP path is both a file and a folder: {name}")
        else:
            if len(raw) > MAX_RECORD_BYTES:
                raise BackupError("JSON records exceed 50 MB.")
            document = json.loads(raw)
    except BackupError:
        raise
    except Exception as exc:
        raise BackupError("The file is damaged or is not a valid JSON, Excel, or ZIP backup.") from exc
    if not isinstance(document, dict) or document.get("schema") != SCHEMA:
        raise BackupError("Unsupported backup version. Use a file exported from Settings.")
    records = document.get("records")
    if not isinstance(records, list) or len(records) > MAX_RECORDS:
        raise BackupError("Backup must contain a records list with at most 100,000 records.")
    if not records and not files:
        raise BackupError("The backup contains no records or files to import.")
    return records, files, hashlib.sha256(raw).hexdigest()


def validate_records(records, user):
    allowed, seen, errors, objects = backup_models(), set(), [], []
    for index, record in enumerate(records, 1):
        location = f"Record {index}"
        try:
            if not isinstance(record, dict) or set(record) != {"model", "pk", "fields"}:
                raise BackupError("expected model, pk, and fields.")
            label = record["model"]
            if not isinstance(label, str) or label not in allowed:
                raise BackupError("unknown or unsupported model.")
            model = allowed[label]
            location = f"Record {index} ({label}, ID {record['pk']})"
            if record["pk"] is None or isinstance(record["pk"], (bool, list, dict)):
                raise BackupError("a valid record ID is required.")
            check_value_type(model._meta.pk, record["pk"])
            pk = model._meta.pk.clean(record["pk"], None)
            if isinstance(model._meta.pk, models.AutoField) and pk < 1:
                raise BackupError("record ID must be positive.")
            key = (label, str(pk))
            if key in seen:
                raise BackupError("duplicate record ID.")
            seen.add(key)
            fields = model_fields(model)
            if not isinstance(record["fields"], dict) or set(record["fields"]) != {f.name for f in fields}:
                raise BackupError("fields must exactly match an exported record for this model.")
            obj = next(serializers.deserialize("python", [record]))
            for field in [model._meta.pk, *fields]:
                value = record["pk"] if field.primary_key else record["fields"][field.name]
                if field.many_to_many:
                    if not isinstance(value, list) or any(isinstance(v, (dict, list, bool)) or v is None for v in value):
                        raise BackupError(f"{field.name}: expected a list of record IDs.")
                    for related_id in value:
                        check_value_type(field.remote_field.model._meta.pk, related_id)
                    continue
                if field.is_relation:
                    if value is None and not field.null:
                        raise BackupError(f"{field.name}: a related record is required.")
                    check_value_type(field.target_field, value)
                    continue  # Database constraints are checked after all records are loaded.
                check_value_type(field, value)
                field.clean(value, obj.object)
                if isinstance(field, models.FileField) and value:
                    media_path(value)
            if label == "auth.user" and pk == user.pk:
                for name in ("username", "password", "is_staff", "is_superuser", "is_active"):
                    if getattr(obj.object, name) != getattr(user, name):
                        raise BackupError("cannot change the signed-in administrator's login or permissions.")
            objects.append(obj)
        except (BackupError, ValidationError, ValueError, TypeError, KeyError, serializers.base.DeserializationError) as exc:
            message = "; ".join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
            errors.append(f"{location}: {message}")
        if len(errors) >= 50:
            errors.append("Only the first 50 errors are shown. Fix them and preview again.")
            break
    if errors:
        raise BackupError(errors)
    # Check links against the final graph, including records later in the file.
    incoming = {(obj.object._meta.label_lower, str(obj.object.pk)) for obj in objects}
    for obj in objects:
        for field in model_fields(type(obj.object)):
            if not field.is_relation:
                continue
            related = field.remote_field.model
            values = (obj.m2m_data or {}).get(field.name, []) if field.many_to_many else [getattr(obj.object, field.attname)]
            for value in values:
                if value is None or (related._meta.label_lower, str(value)) in incoming:
                    continue
                if not related.objects.filter(pk=value).exists():
                    errors.append(f"{obj.object._meta.label_lower}, ID {obj.object.pk}, {field.name}: related ID {value} does not exist.")
                    if len(errors) >= 50:
                        raise BackupError(errors)
    if errors:
        raise BackupError(errors)
    return objects


def save_records(objects):
    # Raw fixture saves preserve timestamps and skip custom save() side effects.
    relations = [(obj.object, obj.m2m_data or {}) for obj in objects]
    for obj in objects:
        try:
            obj.save(save_m2m=False)
        except IntegrityError as exc:
            raise BackupError(
                f"{obj.object._meta.label_lower}, ID {obj.object.pk}: a unique value or database constraint "
                "conflicts with another record. Check unique fields and related IDs. No records were imported."
            ) from exc
    for instance, data in relations:
        for name, values in data.items():
            getattr(instance, name).set(values)
    connection.check_constraints()


def database_fingerprint(files):
    digest = hashlib.sha256(json_bytes(snapshot()["records"]))
    for name in sorted(files):
        target = media_path(name)
        digest.update(name.encode("utf-8"))
        if target.exists():
            with target.open("rb") as source:
                digest.update(hashlib.file_digest(source, "sha256").digest())
        else:
            digest.update(b"missing")
    return digest.hexdigest()


def preview_summary(records, files):
    counts = {}
    missing = set()
    allowed = backup_models()
    for record in records:
        label = record["model"]
        counts.setdefault(label, {"model": label, "create": 0, "update": 0})
        exists = allowed[label].objects.filter(pk=record["pk"]).exists()
        counts[label]["update" if exists else "create"] += 1
        for field in model_fields(allowed[label]):
            name = record["fields"].get(field.name)
            # These profile defaults represent an account without an uploaded
            # avatar/cover. They are not missing user uploads. Keep this scoped
            # to the profile fields so other references still produce warnings.
            if (label in {"accounts.userprofile", "user_api.userprofile"}
                    and field.name in {"avatar", "cover_photo"}
                    and name == field.default
                    and name in {"default.jpg", "default_cover.jpg"}):
                continue
            if isinstance(field, models.FileField) and name and name not in files and not media_path(name).exists():
                missing.add(name)
    replaced = sum(media_path(name).exists() for name in files)
    return {
        "records": len(records), "models": list(counts.values()),
        "create": sum(c["create"] for c in counts.values()),
        "update": sum(c["update"] for c in counts.values()),
        "files": len(files), "replace_files": replaced,
        "file_samples": list(sorted(files))[:10],
        "warnings": ([f"{len(missing)} referenced file(s) are missing from this backup and current uploads: "
                      + ", ".join(sorted(missing)[:10])] if missing else []),
        "samples": [{"model": r["model"], "pk": r["pk"], "fields": {
            k: ("[hidden]" if k == "password" else str(v)[:160])
            for k, v in list(r["fields"].items())[:6]}} for r in records[:5]],
    }


def apply_import(objects, files, *, dry_run=False):
    # Stage both new files and previous copies before making any change. On a
    # database/IO failure, restore files as the database transaction rolls back.
    installed = []
    with tempfile.TemporaryDirectory(prefix="aspin-import-") as temporary:
        staging = Path(temporary)
        if not dry_run:
            for index, (name, content) in enumerate(files.items()):
                target = media_path(name)
                (staging / f"new-{index}").write_bytes(content)
                if target.exists():
                    shutil.copyfile(target, staging / f"old-{index}")
        try:
            with connection.constraint_checks_disabled(), transaction.atomic():
                save_records(objects)
                if dry_run:
                    transaction.set_rollback(True)
                else:
                    for index, name in enumerate(files):
                        target = media_path(name)
                        target.parent.mkdir(parents=True, exist_ok=True)
                        installed.append((index, target))
                        # Temporary sibling ensures atomic replacement on the same volume.
                        with tempfile.NamedTemporaryFile(dir=target.parent, delete=False) as output:
                            sibling = Path(output.name)
                            with (staging / f"new-{index}").open("rb") as source:
                                shutil.copyfileobj(source, output)
                        try:
                            os.replace(sibling, target)
                        finally:
                            sibling.unlink(missing_ok=True)
                    imported_models = list({type(obj.object) for obj in objects})
                    with connection.cursor() as cursor:
                        for sql in connection.ops.sequence_reset_sql(no_style(), imported_models):
                            cursor.execute(sql)
        except Exception:
            for index, target in reversed(installed):
                old = staging / f"old-{index}"
                if old.exists():
                    shutil.copyfile(old, target)
                else:
                    target.unlink(missing_ok=True)
            raise


def export_file(format_name):
    document = snapshot()
    if format_name == "json":
        return io.BytesIO(json_bytes(document)), "application/json"
    if format_name == "xlsx":
        return io.BytesIO(excel_bytes(document)), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if format_name != "zip":
        raise BackupError("Choose json, xlsx, or zip.")
    output = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024)
    try:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("records.json", json_bytes(document))
            root = Path(settings.MEDIA_ROOT).resolve()
            for file in sorted(root.rglob("*")):
                if file.is_file():
                    name = file.relative_to(root).as_posix()
                    archive.write(media_path(name), f"media/{name}")
        output.seek(0)
        return output, "application/zip"
    except Exception:
        output.close()
        raise
